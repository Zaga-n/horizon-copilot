"""Small real PDF and Word documents for parser and pipeline tests."""

from io import BytesIO

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def word_bytes(*, text: str = "Horizon project evidence.", heading: str | None = "Scope") -> bytes:
    document = Document()
    if heading:
        document.add_heading(heading, level=1)
    document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes(*, pages: tuple[str, ...] = ("Horizon project evidence.",)) -> bytes:
    writer = PdfWriter()
    for text in pages:
        page = writer.add_blank_page(width=300, height=300)
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        content = DecodedStreamObject()
        content.set_data(f"BT /F1 12 Tf 10 200 Td ({text}) Tj ET".encode())
        page[NameObject("/Contents")] = content
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()
